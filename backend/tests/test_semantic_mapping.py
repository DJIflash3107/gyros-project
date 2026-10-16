import csv
import io
import json
from unittest.mock import AsyncMock, patch
import openpyxl
import pytest

from app.configs.settings import settings
from app.helpers.exceptions import ValidationException
from app.helpers.error_codes import ErrorCodeRegistry
from app.models.schemas.ai import AIModelInfo
from app.models.sales_report_staging_row import SalesReportStagingRow
from app.services.input_profiler import (
    SpreadsheetProfile,
    profile_excel,
    profile_csv,
    profile_to_prompt_text,
)
from app.services.photo_extractor import extract_from_photo
from app.services.semantic_mapper import (
    CANONICAL_FIELDS,
    ColumnMappingEntry,
    SemanticMappingResult,
    map_columns_with_llm,
)


def _create_test_xlsx(headers: list, rows: list) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _create_test_csv(headers: list, rows: list, delimiter: str = ",") -> bytes:
    sio = io.StringIO()
    writer = csv.writer(sio, delimiter=delimiter)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return sio.getvalue().encode("utf-8")


def _setup_test_data(client, headers):
    biz_resp = client.post("/api/v1/businesses", json={"name": "Semantic Biz"}, headers=headers)
    assert biz_resp.status_code == 201
    biz_id = biz_resp.json()["data"]["business"]["id"]

    branch_resp = client.post(f"/api/v1/businesses/{biz_id}/branches", json={"name": "Main Branch"}, headers=headers)
    assert branch_resp.status_code == 201
    branch_id = branch_resp.json()["data"]["branch"]["id"]

    prod_resp = client.post(f"/api/v1/businesses/{biz_id}/products", json={"name": "Beef Gyros", "base_unit": "portion"}, headers=headers)
    assert prod_resp.status_code == 201
    prod_id = prod_resp.json()["data"]["product"]["id"]

    return biz_id, branch_id, prod_id


# 1. test_profile_excel_basic
def test_profile_excel_basic():
    xlsx_bytes = _create_test_xlsx(
        headers=["Product Name", "Sales Qty", "Sale Date"],
        rows=[
            ["Beef Gyros", 25, "2026-10-01"],
            ["Chicken Gyros", 15, "2026-10-02"],
            ["Lamb Gyros", 10, "2026-10-03"],
        ],
    )
    profile = profile_excel(xlsx_bytes)
    assert profile.headers == ["Product Name", "Sales Qty", "Sale Date"]
    assert profile.total_data_rows == 3
    assert len(profile.sample_rows) == 3
    assert profile.column_types["Product Name"] == "text"
    assert profile.column_types["Sales Qty"] == "numeric"
    assert profile.column_types["Sale Date"] == "date"
    assert profile.has_multi_header is False


# 2. test_profile_csv_basic
def test_profile_csv_basic():
    csv_bytes = _create_test_csv(
        headers=["Item", "Qty", "Price"],
        rows=[
            ["Burger", 10, 50000],
            ["Fries", 5, 20000],
        ],
        delimiter=";",
    )
    profile = profile_csv(csv_bytes)
    assert profile.headers == ["Item", "Qty", "Price"]
    assert profile.total_data_rows == 2
    assert len(profile.sample_rows) == 2
    assert profile.column_types["Item"] == "text"
    assert profile.column_types["Qty"] == "numeric"
    assert profile.column_types["Price"] == "numeric"
    assert profile.merged_cell_ranges == []


# 3. test_profile_merged_cells
def test_profile_merged_cells():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:C1")
    ws["A1"] = "Title Banner"
    ws.append([])  # spacer
    ws.append(["Menu", "Terjual", "Sisa"])
    ws.append(["Kopi", 10, 2])
    bio = io.BytesIO()
    wb.save(bio)

    profile = profile_excel(bio.getvalue())
    assert "A1:C1" in profile.merged_cell_ranges
    assert profile.has_multi_header is True
    assert profile.headers == ["Title Banner", "col_2", "col_3"] or "Menu" in profile.headers


# 4. test_profile_empty_file_raises
def test_profile_empty_file_raises():
    with pytest.raises(ValidationException) as exc_info_xlsx:
        profile_excel(b"")
    assert exc_info_xlsx.value.error_code == ErrorCodeRegistry.RPT_EMPTY_FILE

    with pytest.raises(ValidationException) as exc_info_csv:
        profile_csv(b"")
    assert exc_info_csv.value.error_code == ErrorCodeRegistry.RPT_EMPTY_FILE


# 5. test_llm_mapping_standard_indonesian
@pytest.mark.asyncio
async def test_llm_mapping_standard_indonesian():
    profile = SpreadsheetProfile(
        sheet_names=["Sheet1"],
        active_sheet_name="Sheet1",
        headers=["Nama Barang", "Terjual", "Sisa"],
        sample_rows=[["Kopi Susu", 45, 5]],
        column_types={"Nama Barang": "text", "Terjual": "numeric", "Sisa": "numeric"},
        merged_cell_ranges=[],
        total_data_rows=1,
        has_multi_header=False,
    )

    mock_llm_response = {
        "choices": [
            {
                "message": {
                    "content": json.dumps({
                        "mappings": [
                            {"source_column": "Nama Barang", "target_field": "raw_product_name", "confidence": 0.95, "reasoning": "Product item"},
                            {"source_column": "Terjual", "target_field": "sold_qty", "confidence": 0.98, "reasoning": "Quantity sold"},
                            {"source_column": "Sisa", "target_field": "leftover_qty", "confidence": 0.92, "reasoning": "Remaining stock"},
                        ],
                        "unmapped_columns": [],
                        "warnings": [],
                    })
                }
            }
        ]
    }

    mock_client = AsyncMock()
    mock_client.chat_completion = AsyncMock(return_value=mock_llm_response)

    res = await map_columns_with_llm(profile, client=mock_client)
    assert res.fallback_used is False
    assert len(res.mappings) == 3
    mapping_dict = {m.source_column: m.target_field for m in res.mappings}
    assert mapping_dict["Nama Barang"] == "raw_product_name"
    assert mapping_dict["Terjual"] == "sold_qty"
    assert mapping_dict["Sisa"] == "leftover_qty"
    assert res.mappings[0].confidence == 0.95


# 6. test_llm_mapping_english_headers
@pytest.mark.asyncio
async def test_llm_mapping_english_headers():
    profile = SpreadsheetProfile(
        sheet_names=["Sheet1"],
        active_sheet_name="Sheet1",
        headers=["Product", "Sold", "Revenue"],
        sample_rows=[["Beef Gyros", 10, 500000]],
        column_types={"Product": "text", "Sold": "numeric", "Revenue": "numeric"},
        merged_cell_ranges=[],
        total_data_rows=1,
        has_multi_header=False,
    )

    mock_llm_response = {
        "choices": [
            {
                "message": {
                    "content": json.dumps({
                        "mappings": [
                            {"source_column": "Product", "target_field": "raw_product_name", "confidence": 0.99},
                            {"source_column": "Sold", "target_field": "sold_qty", "confidence": 0.95},
                            {"source_column": "Revenue", "target_field": "recorded_revenue", "confidence": 0.96},
                        ],
                        "unmapped_columns": [],
                        "warnings": [],
                    })
                }
            }
        ]
    }

    mock_client = AsyncMock()
    mock_client.chat_completion = AsyncMock(return_value=mock_llm_response)

    res = await map_columns_with_llm(profile, client=mock_client)
    assert res.fallback_used is False
    target_dict = {m.source_column: m.target_field for m in res.mappings}
    assert target_dict["Product"] == "raw_product_name"
    assert target_dict["Sold"] == "sold_qty"
    assert target_dict["Revenue"] == "recorded_revenue"


# 7. test_llm_mapping_unknown_columns
@pytest.mark.asyncio
async def test_llm_mapping_unknown_columns():
    profile = SpreadsheetProfile(
        sheet_names=["Sheet1"],
        active_sheet_name="Sheet1",
        headers=["Product", "Catatan"],
        sample_rows=[["Gyros", "Pelanggan minta extra saus"]],
        column_types={"Product": "text", "Catatan": "text"},
        merged_cell_ranges=[],
        total_data_rows=1,
        has_multi_header=False,
    )

    mock_llm_response = {
        "choices": [
            {
                "message": {
                    "content": json.dumps({
                        "mappings": [
                            {"source_column": "Product", "target_field": "raw_product_name", "confidence": 0.99},
                            {"source_column": "Catatan", "target_field": "_skip", "confidence": 0.90, "reasoning": "Notes column"},
                        ],
                        "unmapped_columns": ["Catatan"],
                        "warnings": [],
                    })
                }
            }
        ]
    }

    mock_client = AsyncMock()
    mock_client.chat_completion = AsyncMock(return_value=mock_llm_response)

    res = await map_columns_with_llm(profile, client=mock_client)
    catatan_entry = next(m for m in res.mappings if m.source_column == "Catatan")
    assert catatan_entry.target_field == "_skip"
    assert "Catatan" in res.unmapped_columns


# 8. test_llm_mapping_low_confidence_flagged
def test_llm_mapping_low_confidence_flagged(client, auth_user):
    headers = auth_user["headers"]
    biz_id, branch_id, prod_id = _setup_test_data(client, headers)

    xlsx_bytes = _create_test_xlsx(
        headers=["Product", "Branch ID", "Period Start", "Period End", "Granularity", "AmbiguousCol", "Recorded Revenue"],
        rows=[
            ["Beef Gyros", branch_id, "2026-10-01", "2026-10-01", "daily", 15, 750000],
        ],
    )

    low_conf_mapping = {
        "choices": [
            {
                "message": {
                    "content": json.dumps({
                        "mappings": [
                            {"source_column": "Product", "target_field": "raw_product_name", "confidence": 0.95},
                            {"source_column": "Branch ID", "target_field": "branch_id", "confidence": 0.95},
                            {"source_column": "Period Start", "target_field": "period_start", "confidence": 0.95},
                            {"source_column": "Period End", "target_field": "period_end", "confidence": 0.95},
                            {"source_column": "Granularity", "target_field": "granularity", "confidence": 0.95},
                            {"source_column": "AmbiguousCol", "target_field": "sold_qty", "confidence": 0.40, "reasoning": "Unclear header"},
                            {"source_column": "Recorded Revenue", "target_field": "recorded_revenue", "confidence": 0.95},
                        ],
                        "unmapped_columns": [],
                        "warnings": ["Low confidence on AmbiguousCol"],
                    })
                }
            }
        ]
    }

    with patch("app.configs.amd_client.AMDTokenFactoryClient.chat_completion", new_callable=AsyncMock) as mock_chat:
        mock_chat.return_value = low_conf_mapping
        resp = client.post(
            f"/api/v1/businesses/{biz_id}/sales_report_submissions/process",
            data={
                "input_method": "spreadsheet",
                "report_period_type": "daily",
                "coverage_start": "2026-10-01",
                "coverage_end": "2026-10-01",
                "branch_id": branch_id,
            },
            files={"file": ("sales.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )

    assert resp.status_code == 201
    body = resp.json()["data"]
    issues = body["issues"]
    low_conf_issues = [i for i in issues if i["issue_code"] == ErrorCodeRegistry.RPT_LOW_CONFIDENCE]
    assert len(low_conf_issues) >= 1
    assert low_conf_issues[0]["severity"] == "warning"
    assert "AmbiguousCol" in low_conf_issues[0]["message"]


# 9. test_llm_fallback_on_api_error
@pytest.mark.asyncio
async def test_llm_fallback_on_api_error():
    profile = SpreadsheetProfile(
        sheet_names=["Sheet1"],
        active_sheet_name="Sheet1",
        headers=["nama_produk", "terjual"],
        sample_rows=[["Beef Gyros", 10]],
        column_types={"nama_produk": "text", "terjual": "numeric"},
        merged_cell_ranges=[],
        total_data_rows=1,
        has_multi_header=False,
    )

    mock_client = AsyncMock()
    mock_client.chat_completion = AsyncMock(side_effect=Exception("Gateway connection timeout 504"))

    res = await map_columns_with_llm(profile, client=mock_client)
    assert res.fallback_used is True
    assert res.model_used == "static_fallback"
    mapping_dict = {m.source_column: m.target_field for m in res.mappings}
    assert mapping_dict["nama_produk"] == "raw_product_name"
    assert mapping_dict["terjual"] == "sold_qty"


# 10. test_llm_fallback_when_disabled
@pytest.mark.asyncio
async def test_llm_fallback_when_disabled():
    profile = SpreadsheetProfile(
        sheet_names=["Sheet1"],
        active_sheet_name="Sheet1",
        headers=["nama_produk"],
        sample_rows=[["Gyros"]],
        column_types={"nama_produk": "text"},
        merged_cell_ranges=[],
        total_data_rows=1,
        has_multi_header=False,
    )

    with patch.object(settings, "SEMANTIC_MAPPING_ENABLED", False):
        res = await map_columns_with_llm(profile)
        assert res.fallback_used is True
        assert res.model_used == "static_fallback"
        assert res.mappings[0].target_field == "raw_product_name"


# 11. test_llm_mapping_with_product_context
@pytest.mark.asyncio
async def test_llm_mapping_with_product_context():
    profile = SpreadsheetProfile(
        sheet_names=["Sheet1"],
        active_sheet_name="Sheet1",
        headers=["Item", "Sold"],
        sample_rows=[["Gyros Sapi Special", 5]],
        column_types={"Item": "text", "Sold": "numeric"},
        merged_cell_ranges=[],
        total_data_rows=1,
        has_multi_header=False,
    )

    mock_llm_response = {
        "choices": [
            {
                "message": {
                    "content": json.dumps({
                        "mappings": [
                            {"source_column": "Item", "target_field": "raw_product_name", "confidence": 0.95},
                            {"source_column": "Sold", "target_field": "sold_qty", "confidence": 0.95},
                        ],
                        "unmapped_columns": [],
                        "warnings": [],
                    })
                }
            }
        ]
    }

    mock_client = AsyncMock()
    mock_client.chat_completion = AsyncMock(return_value=mock_llm_response)

    biz_context = {"products": [{"name": "Gyros Sapi Special"}, {"name": "Kebab"}]}
    await map_columns_with_llm(profile, business_context=biz_context, client=mock_client)

    # Verify that product context was embedded in the user prompt sent to LLM
    call_args = mock_client.chat_completion.call_args
    messages = call_args.kwargs.get("messages") or call_args[1].get("messages")
    user_msg = next(m["content"] for m in messages if m["role"] == "user")
    assert "Gyros Sapi Special" in user_msg


# 12. test_end_to_end_xlsx_with_llm
def test_end_to_end_xlsx_with_llm(client, auth_user):
    headers = auth_user["headers"]
    biz_id, branch_id, prod_id = _setup_test_data(client, headers)

    xlsx_bytes = _create_test_xlsx(
        headers=["Menu Item", "Toko", "Mulai", "Akhir", "Tipe Waktu", "Jumlah Terjual", "Omzet Penjualan"],
        rows=[
            ["Beef Gyros", branch_id, "2026-10-01", "2026-10-01", "daily", 20, 1000000],
        ],
    )

    mock_llm_response = {
        "choices": [
            {
                "message": {
                    "content": json.dumps({
                        "mappings": [
                            {"source_column": "Menu Item", "target_field": "raw_product_name", "confidence": 0.95},
                            {"source_column": "Toko", "target_field": "branch_id", "confidence": 0.95},
                            {"source_column": "Mulai", "target_field": "period_start", "confidence": 0.95},
                            {"source_column": "Akhir", "target_field": "period_end", "confidence": 0.95},
                            {"source_column": "Tipe Waktu", "target_field": "granularity", "confidence": 0.95},
                            {"source_column": "Jumlah Terjual", "target_field": "sold_qty", "confidence": 0.95},
                            {"source_column": "Omzet Penjualan", "target_field": "recorded_revenue", "confidence": 0.95},
                        ],
                        "unmapped_columns": [],
                        "warnings": [],
                    })
                }
            }
        ]
    }

    with patch("app.configs.amd_client.AMDTokenFactoryClient.chat_completion", new_callable=AsyncMock) as mock_chat:
        mock_chat.return_value = mock_llm_response
        resp = client.post(
            f"/api/v1/businesses/{biz_id}/sales_report_submissions/process",
            data={
                "input_method": "spreadsheet",
                "report_period_type": "daily",
                "coverage_start": "2026-10-01",
                "coverage_end": "2026-10-01",
                "branch_id": branch_id,
            },
            files={"file": ("report.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )

    assert resp.status_code == 201
    body = resp.json()["data"]
    assert body["submission"]["status"] == "reviewed"
    assert body["summary"]["total_rows"] == 1
    assert body["summary"]["ready_rows"] == 1


# 13. test_end_to_end_csv_with_llm
def test_end_to_end_csv_with_llm(client, auth_user):
    headers = auth_user["headers"]
    biz_id, branch_id, prod_id = _setup_test_data(client, headers)

    csv_bytes = _create_test_csv(
        headers=["Barang", "Cabang ID", "Awal", "Akhir", "Granularitas", "Total Terjual", "Pemasukan"],
        rows=[
            ["Beef Gyros", branch_id, "2026-10-01", "2026-10-01", "daily", 12, 600000],
        ],
    )

    mock_llm_response = {
        "choices": [
            {
                "message": {
                    "content": json.dumps({
                        "mappings": [
                            {"source_column": "Barang", "target_field": "raw_product_name", "confidence": 0.95},
                            {"source_column": "Cabang ID", "target_field": "branch_id", "confidence": 0.95},
                            {"source_column": "Awal", "target_field": "period_start", "confidence": 0.95},
                            {"source_column": "Akhir", "target_field": "period_end", "confidence": 0.95},
                            {"source_column": "Granularitas", "target_field": "granularity", "confidence": 0.95},
                            {"source_column": "Total Terjual", "target_field": "sold_qty", "confidence": 0.95},
                            {"source_column": "Pemasukan", "target_field": "recorded_revenue", "confidence": 0.95},
                        ],
                        "unmapped_columns": [],
                        "warnings": [],
                    })
                }
            }
        ]
    }

    with patch("app.configs.amd_client.AMDTokenFactoryClient.chat_completion", new_callable=AsyncMock) as mock_chat:
        mock_chat.return_value = mock_llm_response
        resp = client.post(
            f"/api/v1/businesses/{biz_id}/sales_report_submissions/process",
            data={
                "input_method": "spreadsheet",
                "report_period_type": "daily",
                "coverage_start": "2026-10-01",
                "coverage_end": "2026-10-01",
                "branch_id": branch_id,
            },
            files={"file": ("report.csv", csv_bytes, "text/csv")},
            headers=headers,
        )

    assert resp.status_code == 201
    body = resp.json()["data"]
    assert body["submission"]["status"] == "reviewed"
    assert body["summary"]["total_rows"] == 1
    assert body["summary"]["ready_rows"] == 1


# 14. test_end_to_end_photo_canonical_fields
@pytest.mark.asyncio
async def test_end_to_end_photo_canonical_fields():
    # 1x1 transparent PNG bytes
    fake_png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\rIDATx\x9cc`\x00\x00\x00\x02\x00\x01H\xaf\xa4q\x00\x00\x00\x00IEND\xaeB`\x82"

    mock_vision_models = [
        AIModelInfo(id="Qwen2.5-VL-7B-Instruct", supports_vision=True)
    ]

    mock_chat_response = {
        "choices": [
            {
                "message": {
                    "content": json.dumps({
                        "rows": [
                            {"raw_product_name": "Beef Gyros", "sold_qty": 10, "recorded_revenue": 500000}
                        ]
                    })
                }
            }
        ]
    }

    with patch("app.services.photo_extractor.list_models", new_callable=AsyncMock) as mock_list, \
         patch("app.configs.amd_client.AMDTokenFactoryClient.chat_completion", new_callable=AsyncMock) as mock_chat:
        mock_list.return_value = mock_vision_models
        mock_chat.return_value = mock_chat_response

        rows, meta = await extract_from_photo(
            image_bytes=fake_png,
            mime_type="image/png",
            business_context={"products": [{"name": "Beef Gyros"}]},
        )

        assert len(rows) == 1
        assert rows[0]["raw_product_name"] == "Beef Gyros"

        # Verify prompt contained canonical fields from semantic_mapper.py
        call_args = mock_chat.call_args
        messages = call_args.kwargs.get("messages") or call_args[1].get("messages")
        system_msg = next(m["content"] for m in messages if m["role"] == "system")

        for canonical_name in ["raw_product_name", "recorded_revenue", "sold_qty", "leftover_qty"]:
            assert canonical_name in system_msg


# 15. test_mapping_metadata_persisted
def test_mapping_metadata_persisted(client, auth_user, db_session):
    headers = auth_user["headers"]
    biz_id, branch_id, prod_id = _setup_test_data(client, headers)

    xlsx_bytes = _create_test_xlsx(
        headers=["Product", "Branch ID", "Period Start", "Period End", "Granularity", "Sold Qty", "Recorded Revenue"],
        rows=[
            ["Beef Gyros", branch_id, "2026-10-01", "2026-10-01", "daily", 25, 1250000],
        ],
    )

    mock_llm_response = {
        "choices": [
            {
                "message": {
                    "content": json.dumps({
                        "mappings": [
                            {"source_column": "Product", "target_field": "raw_product_name", "confidence": 0.99, "reasoning": "Standard product column"},
                            {"source_column": "Branch ID", "target_field": "branch_id", "confidence": 0.99},
                            {"source_column": "Period Start", "target_field": "period_start", "confidence": 0.99},
                            {"source_column": "Period End", "target_field": "period_end", "confidence": 0.99},
                            {"source_column": "Granularity", "target_field": "granularity", "confidence": 0.99},
                            {"source_column": "Sold Qty", "target_field": "sold_qty", "confidence": 0.99},
                            {"source_column": "Recorded Revenue", "target_field": "recorded_revenue", "confidence": 0.99},
                        ],
                        "unmapped_columns": [],
                        "warnings": [],
                    })
                }
            }
        ]
    }

    with patch("app.configs.amd_client.AMDTokenFactoryClient.chat_completion", new_callable=AsyncMock) as mock_chat:
        mock_chat.return_value = mock_llm_response
        resp = client.post(
            f"/api/v1/businesses/{biz_id}/sales_report_submissions/process",
            data={
                "input_method": "spreadsheet",
                "report_period_type": "daily",
                "coverage_start": "2026-10-01",
                "coverage_end": "2026-10-01",
                "branch_id": branch_id,
            },
            files={"file": ("sales.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )

    assert resp.status_code == 201
    sub_id = resp.json()["data"]["submission"]["id"]

    # Query staging row from database directly
    row = db_session.query(SalesReportStagingRow).filter(SalesReportStagingRow.submission_id == sub_id).first()
    assert row is not None
    assert row.extraction_confidence_json is not None
    conf = row.extraction_confidence_json
    assert conf["mapping_method"] == "llm"
    assert "column_confidences" in conf
    assert conf["column_confidences"]["Product"]["target"] == "raw_product_name"
    assert conf["column_confidences"]["Product"]["confidence"] == 0.99
