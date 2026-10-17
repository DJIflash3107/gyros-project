from datetime import date
from decimal import Decimal
import io
import json
from unittest.mock import AsyncMock, patch
import openpyxl
import pytest

from app.configs.settings import settings
from app.models.expense import Expense
from app.models.branch import Branch
from app.models.business import Business
from app.services.expense_spreadsheet_parser import (
    parse_currency_amount,
    parse_flexible_date,
    parse_expense_excel,
    parse_expense_csv,
)
from app.services.expense_import_service import normalize_category
from app.services.input_profiler import profile_excel, profile_csv
from app.services.semantic_mapper import (
    map_expense_columns_with_llm,
    build_fallback_expense_mapping,
    categorize_expenses_with_llm,
    CANONICAL_EXPENSE_FIELDS,
)


def _create_expense_xlsx(headers, rows, sheet_name="Sheet1") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for r in rows:
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _create_multi_sheet_xlsx(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    for idx, (name, (headers, rows)) in enumerate(sheets.items()):
        ws = wb.active if idx == 0 else wb.create_sheet(title=name)
        if idx == 0:
            ws.title = name
        if headers:
            ws.append(headers)
        for row in rows:
            ws.append(row)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# 1. Unit Tests for Parsers
def test_parse_currency_amount():
    assert parse_currency_amount(1500000) == Decimal("1500000")
    assert parse_currency_amount(1500000.5) == Decimal("1500000.5")
    assert parse_currency_amount("1500000") == Decimal("1500000")
    assert parse_currency_amount("Rp 1.500.000") == Decimal("1500000")
    assert parse_currency_amount("Rp. 1.500.000,00") == Decimal("1500000.00")
    assert parse_currency_amount("1,500,000.50") == Decimal("1500000.50")
    assert parse_currency_amount("(50.000)") == Decimal("50000")
    assert parse_currency_amount("Rp 75.000") == Decimal("75000")
    assert parse_currency_amount("-") is None
    assert parse_currency_amount("") is None
    assert parse_currency_amount(None) is None


def test_parse_flexible_date():
    assert parse_flexible_date(date(2026, 10, 15)) == date(2026, 10, 15)
    assert parse_flexible_date("2026-10-15") == date(2026, 10, 15)
    assert parse_flexible_date("15/10/2026") == date(2026, 10, 15)
    assert parse_flexible_date("15-10-2026") == date(2026, 10, 15)
    assert parse_flexible_date("2026/10/15") == date(2026, 10, 15)
    # Excel serial number ~45210 -> Oct 2023
    excel_date = parse_flexible_date(45210)
    assert excel_date is not None
    assert excel_date.year == 2023
    assert parse_flexible_date("invalid_date", default_date=date(2026, 1, 1)) == date(2026, 1, 1)


def test_normalize_category():
    assert normalize_category("Bahan Baku") == "Ingredients"
    assert normalize_category("ingredients") == "Ingredients"
    assert normalize_category("utilities") == "Utilities"
    assert normalize_category("gaji") == "Salaries & Wages"
    assert normalize_category(None, "Beli daging sapi 5kg dan bumbu") == "Ingredients"
    assert normalize_category(None, "Beli kantong kresek dan box gyro") == "Packaging"
    assert normalize_category(None, "Token listrik PLN cabang") == "Utilities"
    assert normalize_category(None, "Servis chiller dan freezer") == "Maintenance"
    assert normalize_category(None, "Cetak spanduk promo oktober") == "Marketing"
    assert normalize_category(None, "Beli tabung gas LPG 12kg") == "Operations"


# 2. Integration Tests: Spreadsheet Import (Excel & CSV)
def test_import_expenses_excel_preview_and_persist(client, auth_user, db_session, monkeypatch):
    monkeypatch.setattr(settings, "SEMANTIC_MAPPING_ENABLED", False)
    headers = auth_user["headers"]

    # Setup Business & Branch
    biz_resp = client.post("/api/v1/businesses", json={"name": "Expense Import Biz"}, headers=headers)
    biz_id = biz_resp.json()["data"]["business"]["id"]

    branch_resp = client.post(f"/api/v1/businesses/{biz_id}/branches", json={"name": "Main Branch"}, headers=headers)
    branch_id = branch_resp.json()["data"]["branch"]["id"]

    xlsx_bytes = _create_expense_xlsx(
        headers=["Tanggal", "Kategori", "Keterangan", "Nominal", "Cabang"],
        rows=[
            ["2026-10-01", "Bahan Baku", "Beli Daging Sapi 10kg", "1.200.000", "Main Branch"],
            ["2026-10-02", "Operasional", "Beli Gas LPG 12kg", "225.000", "Main Branch"],
            ["2026-10-03", "", "Token Listrik PLN", "500000", ""],
        ],
        sheet_name="Pengeluaran",
    )

    # 1. Preview Mode (persist=false)
    resp_preview = client.post(
        f"/api/v1/businesses/{biz_id}/expenses/import",
        data={
            "input_method": "spreadsheet",
            "branch_id": branch_id,
            "persist": "false",
        },
        files={"file": ("expenses.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert resp_preview.status_code == 200
    body_preview = resp_preview.json()["data"]
    assert body_preview["summary"]["total_extracted"] == 3
    assert body_preview["summary"]["total_saved"] == 0
    assert body_preview["summary"]["persisted"] is False
    assert len(body_preview["items"]) == 3
    assert body_preview["items"][2]["category"] == "Utilities"  # Inferred from "Token Listrik PLN"

    # Verify nothing was persisted yet
    assert db_session.query(Expense).filter(Expense.business_id == biz_id).count() == 0

    # 2. Persist Mode (persist=true)
    resp_persist = client.post(
        f"/api/v1/businesses/{biz_id}/expenses/import",
        data={
            "input_method": "spreadsheet",
            "branch_id": branch_id,
            "persist": "true",
        },
        files={"file": ("expenses.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert resp_persist.status_code == 201
    body_persist = resp_persist.json()["data"]
    assert body_persist["summary"]["total_extracted"] == 3
    assert body_persist["summary"]["total_saved"] == 3
    assert body_persist["summary"]["persisted"] is True
    assert float(body_persist["summary"]["total_amount"]) == 1925000.0
    assert len(body_persist["expenses"]) == 3

    # Verify database records
    db_expenses = db_session.query(Expense).filter(Expense.business_id == biz_id).all()
    assert len(db_expenses) == 3


def test_import_expenses_csv(client, auth_user, db_session, monkeypatch):
    monkeypatch.setattr(settings, "SEMANTIC_MAPPING_ENABLED", False)
    headers = auth_user["headers"]
    biz_resp = client.post("/api/v1/businesses", json={"name": "CSV Expense Biz"}, headers=headers)
    biz_id = biz_resp.json()["data"]["business"]["id"]

    csv_content = (
        "Tanggal,Kategori,Deskripsi,Total,Cabang\n"
        "2026-10-10,Kemasan,Kotak Gyro 500 pcs,350000,\n"
        "2026-10-11,Operasional,Es Batu Kristal 10 karung,120000,\n"
    ).encode("utf-8")

    resp = client.post(
        f"/api/v1/businesses/{biz_id}/expenses/import",
        data={
            "input_method": "spreadsheet",
            "persist": "true",
        },
        files={"file": ("expenses.csv", csv_content, "text/csv")},
        headers=headers,
    )
    assert resp.status_code == 201
    body = resp.json()["data"]
    assert body["summary"]["total_saved"] == 2
    assert float(body["summary"]["total_amount"]) == 470000.0


# 3. Integration Tests: Photo OCR Receipt Import
@pytest.mark.asyncio
async def test_import_expenses_photo_ocr(client, auth_user, db_session, monkeypatch):
    headers = auth_user["headers"]
    biz_resp = client.post("/api/v1/businesses", json={"name": "Photo Expense Biz"}, headers=headers)
    biz_id = biz_resp.json()["data"]["business"]["id"]

    mock_llm_json = {
        "summary": {
            "merchant_name": "Toko Sumber Rejeki",
            "transaction_date": "2026-10-16",
            "total_amount": 750000,
            "currency": "IDR",
            "suggested_category": "Ingredients",
            "description": "Pembelian Beras 25kg dan Minyak Goreng 10L",
        },
        "items": [
            {"item_name": "Beras Premium 25kg", "qty": 1, "price": 350000, "subtotal": 350000, "category": "Ingredients"},
            {"item_name": "Minyak Goreng 10L", "qty": 1, "price": 400000, "subtotal": 400000, "category": "Ingredients"},
        ],
    }

    class MockClient:
        async def chat_completion(self, *args, **kwargs):
            return {
                "choices": [
                    {"message": {"content": json.dumps(mock_llm_json)}}
                ]
            }

    monkeypatch.setattr(settings, "AMD_API_KEY", "test-api-key")
    monkeypatch.setattr(settings, "VISION_MODEL_OVERRIDE", "test-vision-model")

    with patch("app.services.expense_photo_extractor.get_amd_client", return_value=MockClient()):
        fake_img = b"fake_jpeg_bytes_for_receipt"
        resp = client.post(
            f"/api/v1/businesses/{biz_id}/expenses/import",
            data={
                "input_method": "photo",
                "persist": "true",
            },
            files={"file": ("receipt.jpg", fake_img, "image/jpeg")},
            headers=headers,
        )
        assert resp.status_code == 201
        body = resp.json()["data"]
        assert body["summary"]["total_saved"] == 1
        assert float(body["summary"]["total_amount"]) == 750000.0
        assert body["expenses"][0]["category"] == "Ingredients"
        assert "Toko Sumber Rejeki" in body["expenses"][0]["description"]


# 4. Multi-Sheet Workbook with Sales and Expenses Integration
def test_multi_sheet_workbook_with_expenses(client, auth_user, db_session, monkeypatch):
    monkeypatch.setattr(settings, "SEMANTIC_MAPPING_ENABLED", False)
    headers = auth_user["headers"]

    biz_resp = client.post("/api/v1/businesses", json={"name": "MultiSheet With Expenses Biz"}, headers=headers)
    biz_id = biz_resp.json()["data"]["business"]["id"]

    branch_resp = client.post(f"/api/v1/businesses/{biz_id}/branches", json={"name": "Central Branch"}, headers=headers)
    branch_id = branch_resp.json()["data"]["branch"]["id"]

    client.post(f"/api/v1/businesses/{biz_id}/products", json={"name": "Beef Gyros", "base_unit": "portion"}, headers=headers)

    xlsx_bytes = _create_multi_sheet_xlsx({
        "Penjualan": (
            ["Product", "Branch ID", "Period Start", "Period End", "Granularity", "Sold Qty", "Recorded Revenue"],
            [["Beef Gyros", branch_id, "2026-10-01", "2026-10-01", "daily", 15, 750000]],
        ),
        "Biaya Operasional": (
            ["Tanggal", "Kategori", "Keterangan", "Nominal"],
            [
                ["2026-10-01", "Bahan Baku", "Daging Sapi", "300000"],
                ["2026-10-01", "", "Listrik Token", "150000"],
            ],
        ),
    })

    resp = client.post(
        f"/api/v1/businesses/{biz_id}/sales_report_submissions/process",
        data={
            "input_method": "spreadsheet",
            "report_period_type": "daily",
            "coverage_start": "2026-10-01",
            "coverage_end": "2026-10-01",
            "branch_id": branch_id,
        },
        files={"file": ("sales_and_expenses.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert resp.status_code == 201
    body = resp.json()["data"]
    sub_id = body["submission"]["id"]

    # Verify sales staging rows
    assert body["summary"]["total_rows"] == 1
    assert "Penjualan" in body["summary"]["sheets_processed"]
    assert "Biaya Operasional" in body["summary"]["sheets_processed"]

    # Verify extracted expenses metrics
    assert body["summary"]["expenses_extracted"] == 2
    assert float(body["summary"]["expenses_total_amount"]) == 450000.0

    # Verify expenses are saved in database linked to source_submission_id
    expenses = db_session.query(Expense).filter(Expense.source_submission_id == sub_id).all()
    assert len(expenses) == 2
    exp_cats = {e.category for e in expenses}
    assert "Ingredients" in exp_cats
    assert "Utilities" in exp_cats  # "Listrik Token" auto-mapped to Utilities


# 5. Error Handling Tests
def test_import_expenses_no_file(client, auth_user):
    headers = auth_user["headers"]
    biz_resp = client.post("/api/v1/businesses", json={"name": "Err Biz"}, headers=headers)
    biz_id = biz_resp.json()["data"]["business"]["id"]

    resp = client.post(
        f"/api/v1/businesses/{biz_id}/expenses/import",
        data={"input_method": "spreadsheet"},
        headers=headers,
    )
    assert resp.status_code == 422


# 6. Semantic Mapping & AI Categorization Tests
@pytest.mark.asyncio
async def test_map_expense_columns_with_llm_primary(monkeypatch):
    monkeypatch.setattr(settings, "SEMANTIC_MAPPING_ENABLED", True)
    monkeypatch.setattr(settings, "AMD_API_KEY", "test-key")

    mock_llm_mapping = {
        "mappings": [
            {"source_column": "Tgl Transaksi", "target_field": "expense_date", "confidence": 0.98, "reasoning": "Transaction date"},
            {"source_column": "Pos Akun", "target_field": "category", "confidence": 0.92, "reasoning": "Accounting category"},
            {"source_column": "Uraian Pembayaran", "target_field": "description", "confidence": 0.95, "reasoning": "Payment description"},
            {"source_column": "Nominal Keluar", "target_field": "amount", "confidence": 0.99, "reasoning": "Expense amount"},
            {"source_column": "Gerai", "target_field": "branch_id", "confidence": 0.90, "reasoning": "Branch outlet"},
            {"source_column": "No Ref", "target_field": "_skip", "confidence": 0.85, "reasoning": "Reference number not needed"},
        ],
        "unmapped_columns": [],
        "warnings": [],
    }

    class MockClient:
        async def chat_completion(self, *args, **kwargs):
            return {
                "choices": [
                    {"message": {"content": json.dumps(mock_llm_mapping)}}
                ]
            }
        async def close(self):
            pass

    xlsx_bytes = _create_expense_xlsx(
        headers=["Tgl Transaksi", "Pos Akun", "Uraian Pembayaran", "Nominal Keluar", "Gerai", "No Ref"],
        rows=[["2026-10-01", "Bahan Baku", "Daging Sapi", "500000", "Main Branch", "REF123"]],
        sheet_name="Pengeluaran",
    )
    profile = profile_excel(xlsx_bytes)

    result = await map_expense_columns_with_llm(profile, client=MockClient())

    assert result.fallback_used is False
    assert len(result.mappings) >= 6
    map_dict = {m.source_column: m.target_field for m in result.mappings}
    assert map_dict["Tgl Transaksi"] == "expense_date"
    assert map_dict["Pos Akun"] == "category"
    assert map_dict["Uraian Pembayaran"] == "description"
    assert map_dict["Nominal Keluar"] == "amount"
    assert map_dict["Gerai"] == "branch_id"
    assert map_dict["No Ref"] == "_skip"


@pytest.mark.asyncio
async def test_map_expense_columns_fallback(monkeypatch):
    monkeypatch.setattr(settings, "SEMANTIC_MAPPING_ENABLED", False)

    xlsx_bytes = _create_expense_xlsx(
        headers=["Tanggal", "Kategori", "Keterangan", "Jumlah", "Cabang"],
        rows=[["2026-10-01", "Kemasan", "Box", "200000", "Main Branch"]],
    )
    profile = profile_excel(xlsx_bytes)

    result = await map_expense_columns_with_llm(profile)
    assert result.fallback_used is True
    map_dict = {m.source_column: m.target_field for m in result.mappings}
    assert map_dict["Tanggal"] == "expense_date"
    assert map_dict["Kategori"] == "category"
    assert map_dict["Keterangan"] == "description"
    assert map_dict["Jumlah"] == "amount"
    assert map_dict["Cabang"] == "branch"


@pytest.mark.asyncio
async def test_categorize_expenses_with_llm_primary(monkeypatch):
    monkeypatch.setattr(settings, "SEMANTIC_MAPPING_ENABLED", True)
    monkeypatch.setattr(settings, "AMD_API_KEY", "test-key")

    mock_cats_response = {
        "classifications": [
            {"index": 0, "category": "Ingredients", "confidence": 0.97},
            {"index": 1, "category": "Packaging", "confidence": 0.95},
            {"index": 2, "category": "Utilities", "confidence": 0.99},
        ]
    }

    class MockClient:
        async def chat_completion(self, *args, **kwargs):
            return {
                "choices": [
                    {"message": {"content": json.dumps(mock_cats_response)}}
                ]
            }
        async def close(self):
            pass

    items = [
        {"description": "Daging slice 5kg", "category": ""},
        {"description": "Paper lunch box L 500 pcs", "category": ""},
        {"description": "Listrik prabayar 500k", "category": ""},
    ]

    cats = await categorize_expenses_with_llm(items, client=MockClient())
    assert cats == ["Ingredients", "Packaging", "Utilities"]


def test_end_to_end_expense_import_with_llm(client, auth_user, db_session, monkeypatch):
    monkeypatch.setattr(settings, "SEMANTIC_MAPPING_ENABLED", True)
    monkeypatch.setattr(settings, "AMD_API_KEY", "test-key")

    headers = auth_user["headers"]
    biz_resp = client.post("/api/v1/businesses", json={"name": "LLM Semantic Expense Biz"}, headers=headers)
    biz_id = biz_resp.json()["data"]["business"]["id"]

    mock_llm_mapping = {
        "mappings": [
            {"source_column": "Hari/Tgl", "target_field": "expense_date", "confidence": 0.96},
            {"source_column": "Pos", "target_field": "category", "confidence": 0.90},
            {"source_column": "Deskripsi Belanja", "target_field": "description", "confidence": 0.97},
            {"source_column": "Total Bayar", "target_field": "amount", "confidence": 0.98},
        ],
        "unmapped_columns": [],
        "warnings": [],
    }

    mock_cats_response = {
        "classifications": [
            {"index": 0, "category": "Ingredients", "confidence": 0.95},
            {"index": 1, "category": "Maintenance", "confidence": 0.92},
        ]
    }

    class MockClient:
        async def chat_completion(self, *args, **kwargs):
            messages = kwargs.get("messages", [])
            sys_text = messages[0]["content"] if messages else ""
            if "financial schema mapping" in sys_text:
                return {"choices": [{"message": {"content": json.dumps(mock_llm_mapping)}}]}
            elif "bookkeeping AI specializing" in sys_text:
                return {"choices": [{"message": {"content": json.dumps(mock_cats_response)}}]}
            return {"choices": [{"message": {"content": "{}"}}]}
        async def close(self):
            pass

    xlsx_bytes = _create_expense_xlsx(
        headers=["Hari/Tgl", "Pos", "Deskripsi Belanja", "Total Bayar"],
        rows=[
            ["2026-10-01", "", "Beli saus keju 5 drum", "450000"],
            ["2026-10-02", "", "Ganti freon chiller dapur", "350000"],
        ],
        sheet_name="Pengeluaran Toko",
    )

    with patch("app.services.semantic_mapper.AMDTokenFactoryClient", return_value=MockClient()):
        resp = client.post(
            f"/api/v1/businesses/{biz_id}/expenses/import",
            data={
                "input_method": "spreadsheet",
                "persist": "true",
            },
            files={"file": ("colloquial_expenses.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )
        assert resp.status_code == 201
        body = resp.json()["data"]
        assert body["summary"]["total_saved"] == 2
        assert float(body["summary"]["total_amount"]) == 800000.0
        exp_list = body["expenses"]
        assert exp_list[0]["category"] == "Ingredients"
        assert exp_list[1]["category"] == "Maintenance"


def test_preview_correct_and_batch_save_workflow(client, auth_user, db_session, monkeypatch):
    monkeypatch.setattr(settings, "SEMANTIC_MAPPING_ENABLED", False)
    headers = auth_user["headers"]
    biz_resp = client.post("/api/v1/businesses", json={"name": "Preview and Batch Biz"}, headers=headers)
    biz_id = biz_resp.json()["data"]["business"]["id"]

    branch_resp = client.post(f"/api/v1/businesses/{biz_id}/branches", json={"name": "Outlet Alpha"}, headers=headers)
    branch_id = branch_resp.json()["data"]["branch"]["id"]

    xlsx_bytes = _create_expense_xlsx(
        headers=["Tanggal", "Kategori", "Keterangan", "Nominal"],
        rows=[
            ["2026-10-01", "Bahan Baku", "Daging Ayam 10kg", "350000"],
            ["2026-10-02", "Operasional", "Biaya Tak Terduga", "120000"],
        ],
    )

    # 1. Preview mode (persist=false)
    resp_preview = client.post(
        f"/api/v1/businesses/{biz_id}/expenses/import",
        data={
            "input_method": "spreadsheet",
            "branch_id": branch_id,
            "persist": "false",
        },
        files={"file": ("expenses.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert resp_preview.status_code == 200
    preview_items = resp_preview.json()["data"]["items"]
    assert len(preview_items) == 2
    assert db_session.query(Expense).filter(Expense.business_id == biz_id).count() == 0

    # 2. Simulate User correcting the second item's category and description
    corrected_items = [
        {
            "branch_id": preview_items[0]["branch_id"],
            "expense_date": preview_items[0]["expense_date"],
            "category": preview_items[0]["category"],
            "description": preview_items[0]["description"],
            "amount": preview_items[0]["amount"],
            "currency": preview_items[0]["currency"],
        },
        {
            "branch_id": branch_id,
            "expense_date": "2026-10-02",
            "category": "Maintenance",  # Corrected from Operations
            "description": "Perbaikan Pintu Gudang Toko",  # Corrected description
            "amount": "120000.00",
            "currency": "IDR",
        },
    ]

    # 3. Save corrected batch directly via POST /expenses/batch
    resp_batch = client.post(
        f"/api/v1/businesses/{biz_id}/expenses/batch",
        json={"items": corrected_items},
        headers=headers,
    )
    assert resp_batch.status_code == 201
    batch_data = resp_batch.json()["data"]
    assert batch_data["summary"]["total_saved"] == 2
    assert float(batch_data["summary"]["total_amount"]) == 470000.0

    # 4. Verify in DB
    db_items = db_session.query(Expense).filter(Expense.business_id == biz_id).order_by(Expense.id.asc()).all()
    assert len(db_items) == 2
    assert db_items[0].category == "Ingredients"
    assert db_items[1].category == "Maintenance"
    assert db_items[1].description == "Perbaikan Pintu Gudang Toko"



