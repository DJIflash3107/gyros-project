from dataclasses import dataclass
from typing import List, Dict, Any, Optional
import io
import csv
import re
from datetime import date, datetime
from decimal import Decimal
import openpyxl

from app.helpers.error_codes import ErrorCodeRegistry
from app.helpers.exceptions import ValidationException
from app.helpers.logging import logger


@dataclass
class SpreadsheetProfile:
    """Structured profile of a spreadsheet's layout for LLM consumption."""
    sheet_names: List[str]
    active_sheet_name: str
    headers: List[str]               # Raw original header strings
    sample_rows: List[List[Any]]     # First 5 non-empty data rows
    column_types: Dict[str, str]     # Inferred type per header: "numeric", "date", "text", "mixed"
    merged_cell_ranges: List[str]    # e.g. ["A1:C1"] for merged cells
    total_data_rows: int
    has_multi_header: bool           # True if headers span >1 row


def infer_column_type(values: List[Any]) -> str:
    """Infer data type for a list of values."""
    types_found = set()
    for v in values:
        if v is None:
            continue
        s = str(v).strip()
        if not s:
            continue

        if isinstance(v, (datetime, date)):
            types_found.add("date")
            continue

        if isinstance(v, (int, float, Decimal)):
            types_found.add("numeric")
            continue

        # Try parsing string as number (handling commas/currency)
        clean_s = re.sub(r"^[^\d\-+]+|[^\d\.]+$", "", s).replace(",", "")
        try:
            float(clean_s)
            types_found.add("numeric")
            continue
        except (ValueError, TypeError):
            pass

        # Try parsing string as date
        is_date = False
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                datetime.strptime(s, fmt)
                is_date = True
                break
            except ValueError:
                pass
        if is_date:
            types_found.add("date")
            continue

        types_found.add("text")

    if not types_found:
        return "text"
    if len(types_found) == 1:
        return types_found.pop()
    return "mixed"


def profile_to_prompt_text(profile: SpreadsheetProfile) -> str:
    """Convert a SpreadsheetProfile into a concise prompt text for LLM mapping."""
    lines = [
        f"Sheet Names: {', '.join(profile.sheet_names)}",
        f"Active Sheet: {profile.active_sheet_name}",
        f"Total Data Rows: {profile.total_data_rows}",
        f"Multi-Header Detected: {'Yes' if profile.has_multi_header else 'No'}",
    ]
    if profile.merged_cell_ranges:
        lines.append(f"Merged Cell Ranges: {', '.join(profile.merged_cell_ranges)}")
    else:
        lines.append("Merged Cell Ranges: None")

    lines.append("\nDetected Columns & Types:")
    for idx, h in enumerate(profile.headers, start=1):
        ctype = profile.column_types.get(h, "text")
        lines.append(f"  {idx}. \"{h}\" (inferred type: {ctype})")

    lines.append("\nSample Data Rows (up to 5):")
    if not profile.sample_rows:
        lines.append("  (No sample data available)")
    else:
        for idx, row in enumerate(profile.sample_rows, start=1):
            formatted_row = [str(val) if val is not None else "" for val in row]
            lines.append(f"  Row {idx}: {formatted_row}")

    return "\n".join(lines)


def profile_excel(file_bytes: bytes) -> SpreadsheetProfile:
    """Profile an Excel file's structure, headers, sample rows, and data types."""
    if not file_bytes:
        raise ValidationException(
            errors={"file": "Empty spreadsheet file", "error_code": ErrorCodeRegistry.RPT_EMPTY_FILE},
            message="Spreadsheet contains no data rows",
            error_code=ErrorCodeRegistry.RPT_EMPTY_FILE,
        )

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        logger.error(f"Error loading Excel workbook for profiling: {e}", exc_info=True)
        raise ValidationException(
            errors={"file": f"Invalid or corrupt Excel workbook: {str(e)}", "error_code": ErrorCodeRegistry.RPT_MALFORMED_FILE},
            message="Malformed or corrupt spreadsheet file",
            error_code=ErrorCodeRegistry.RPT_MALFORMED_FILE,
        )

    sheet_names = wb.sheetnames
    if not sheet_names:
        raise ValidationException(
            errors={"file": "Workbook contains no worksheets", "error_code": ErrorCodeRegistry.RPT_EMPTY_FILE},
            message="Spreadsheet contains no data rows",
            error_code=ErrorCodeRegistry.RPT_EMPTY_FILE,
        )

    sheet = wb.active
    if sheet is None:
        sheet = wb[sheet_names[0]]

    active_sheet_name = sheet.title
    merged_ranges = [str(r) for r in sheet.merged_cells.ranges]

    has_multi_header = False
    for r in sheet.merged_cells.ranges:
        if r.min_row <= 2:
            has_multi_header = True
            break

    rows_iter = sheet.iter_rows(values_only=True)
    header_row = None
    for r in rows_iter:
        if any(cell is not None and str(cell).strip() != "" for cell in r):
            header_row = r
            break

    if not header_row:
        raise ValidationException(
            errors={"file": "Workbook contains no data or headers", "error_code": ErrorCodeRegistry.RPT_EMPTY_FILE},
            message="Spreadsheet contains no data rows",
            error_code=ErrorCodeRegistry.RPT_EMPTY_FILE,
        )

    headers = [str(c).strip() if c is not None else f"col_{i+1}" for i, c in enumerate(header_row)]
    sample_rows = []
    total_data_rows = 0
    col_values: Dict[int, List[Any]] = {i: [] for i in range(len(headers))}

    for r in rows_iter:
        if not any(cell is not None and str(cell).strip() != "" for cell in r):
            continue
        total_data_rows += 1
        row_vals = list(r)[:len(headers)]
        if len(row_vals) < len(headers):
            row_vals.extend([None] * (len(headers) - len(row_vals)))

        if len(sample_rows) < 5:
            sample_rows.append(row_vals)

        for i, val in enumerate(row_vals):
            col_values[i].append(val)

    if total_data_rows == 0:
        raise ValidationException(
            errors={"file": "Spreadsheet contains headers but no data rows", "error_code": ErrorCodeRegistry.RPT_EMPTY_FILE},
            message="Spreadsheet contains no data rows",
            error_code=ErrorCodeRegistry.RPT_EMPTY_FILE,
        )

    column_types = {}
    for i, h in enumerate(headers):
        column_types[h] = infer_column_type(col_values.get(i, []))

    return SpreadsheetProfile(
        sheet_names=sheet_names,
        active_sheet_name=active_sheet_name,
        headers=headers,
        sample_rows=sample_rows,
        column_types=column_types,
        merged_cell_ranges=merged_ranges,
        total_data_rows=total_data_rows,
        has_multi_header=has_multi_header,
    )


def profile_csv(file_bytes: bytes) -> SpreadsheetProfile:
    """Profile a CSV file's structure, headers, sample rows, and data types."""
    if not file_bytes:
        raise ValidationException(
            errors={"file": "Empty CSV file", "error_code": ErrorCodeRegistry.RPT_EMPTY_FILE},
            message="Spreadsheet contains no data rows",
            error_code=ErrorCodeRegistry.RPT_EMPTY_FILE,
        )

    text_content = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text_content = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if text_content is None:
        raise ValidationException(
            errors={"file": "Failed to decode CSV text with supported encodings", "error_code": ErrorCodeRegistry.RPT_MALFORMED_FILE},
            message="Malformed or corrupt CSV file",
            error_code=ErrorCodeRegistry.RPT_MALFORMED_FILE,
        )

    if not text_content.strip():
        raise ValidationException(
            errors={"file": "Spreadsheet contains no data rows", "error_code": ErrorCodeRegistry.RPT_EMPTY_FILE},
            message="Spreadsheet contains no data rows",
            error_code=ErrorCodeRegistry.RPT_EMPTY_FILE,
        )

    stream = io.StringIO(text_content)
    sample = text_content[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except Exception:
        dialect = csv.excel

    reader = csv.reader(stream, dialect=dialect)
    header_row = None
    for r in reader:
        if any(cell.strip() for cell in r):
            header_row = r
            break

    if not header_row:
        raise ValidationException(
            errors={"file": "CSV contains no header row", "error_code": ErrorCodeRegistry.RPT_EMPTY_FILE},
            message="Spreadsheet contains no data rows",
            error_code=ErrorCodeRegistry.RPT_EMPTY_FILE,
        )

    headers = [str(c).strip() if c.strip() else f"col_{i+1}" for i, c in enumerate(header_row)]
    sample_rows = []
    total_data_rows = 0
    col_values: Dict[int, List[Any]] = {i: [] for i in range(len(headers))}

    for r in reader:
        if not any(cell.strip() for cell in r):
            continue
        total_data_rows += 1
        row_vals = list(r)[:len(headers)]
        if len(row_vals) < len(headers):
            row_vals.extend([""] * (len(headers) - len(row_vals)))

        if len(sample_rows) < 5:
            sample_rows.append(row_vals)

        for i, val in enumerate(row_vals):
            col_values[i].append(val)

    if total_data_rows == 0:
        raise ValidationException(
            errors={"file": "CSV contains headers but no data rows", "error_code": ErrorCodeRegistry.RPT_EMPTY_FILE},
            message="Spreadsheet contains no data rows",
            error_code=ErrorCodeRegistry.RPT_EMPTY_FILE,
        )

    column_types = {}
    for i, h in enumerate(headers):
        column_types[h] = infer_column_type(col_values.get(i, []))

    return SpreadsheetProfile(
        sheet_names=["Sheet1"],
        active_sheet_name="Sheet1",
        headers=headers,
        sample_rows=sample_rows,
        column_types=column_types,
        merged_cell_ranges=[],
        total_data_rows=total_data_rows,
        has_multi_header=False,
    )
