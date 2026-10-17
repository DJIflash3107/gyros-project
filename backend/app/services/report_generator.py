import csv
from datetime import datetime, date
from decimal import Decimal
import io
from typing import List, Dict, Any, Tuple, Optional

from sqlalchemy import select, func, and_
from sqlalchemy.orm import Session

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.models.artifact import Artifact
from app.models.sales_report_observation import SalesReportObservation
from app.models.product import Product
from app.models.branch import Branch
from app.models.expense import Expense
from app.models.forecast_run import ForecastRun
from app.models.forecast_result import ForecastResult
from app.models.business_action import BusinessAction
from app.models.promotion_plan import PromotionPlan
from app.models.business import Business
from app.models.schemas.artifact import GenerateReportRequest, GeneratePlanArtifactRequest
from app.helpers.exceptions import (
    AppException,
    NotFoundException,
    ValidationException,
    ArtifactGenerationException,
)
from app.helpers.error_codes import ErrorCodeRegistry


MIME_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


class ReportGeneratorService:
    """Service for generating business reports and structured plan exports in CSV, XLSX, and PDF."""

    # ---------------------------------------------------------
    # Format Exporters
    # ---------------------------------------------------------

    @staticmethod
    def _export_csv(headers: List[str], rows: List[List[Any]]) -> bytes:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([str(c) if c is not None else "" for c in row])
        # UTF-8 BOM so Excel opens it with proper encoding
        return output.getvalue().encode("utf-8-sig")

    @staticmethod
    def _export_xlsx(
        sheet_title: str,
        headers: List[str],
        rows: List[List[Any]],
        summary_rows: Optional[List[List[Any]]] = None,
    ) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_title[:31]  # Excel max sheet title is 31 chars

        # Header styling
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )
        data_font = Font(name="Calibri", size=10)

        # Write headers
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Write data rows
        for row_idx, row_data in enumerate(rows, start=2):
            ws.append([r if r is not None else "" for r in row_data])
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if isinstance(row_data[col_idx - 1], (int, float, Decimal)):
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(horizontal="left")

        # Summary rows
        if summary_rows:
            sum_font = Font(name="Calibri", size=10, bold=True)
            sum_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            for sum_data in summary_rows:
                ws.append([r if r is not None else "" for r in sum_data])
                cur_row = ws.max_row
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=cur_row, column=col_idx)
                    cell.font = sum_font
                    cell.fill = sum_fill
                    cell.border = thin_border

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def _export_pdf(
        title: str,
        subtitle: str,
        headers: List[str],
        rows: List[List[Any]],
        summary_text: Optional[str] = None,
        landscape_orientation: bool = False,
    ) -> bytes:
        buffer = io.BytesIO()
        page_size = landscape(letter) if landscape_orientation else letter
        doc = SimpleDocTemplate(
            buffer,
            pagesize=page_size,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36,
        )
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1F4E79"),
            spaceAfter=6,
        )
        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#555555"),
            spaceAfter=14,
        )
        cell_style = ParagraphStyle(
            "TableCell",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
        )
        header_cell_style = ParagraphStyle(
            "HeaderCell",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.white,
            fontName="Helvetica-Bold",
        )

        elements = []
        elements.append(Paragraph(title, title_style))
        elements.append(Paragraph(subtitle, subtitle_style))

        if summary_text:
            summary_style = ParagraphStyle(
                "SummaryText",
                parent=styles["Normal"],
                fontSize=9,
                leading=12,
                textColor=colors.HexColor("#333333"),
                spaceAfter=10,
            )
            elements.append(Paragraph(summary_text, summary_style))
            elements.append(Spacer(1, 0.1 * inch))

        # Format table data
        table_data = []
        # Header row
        table_data.append([Paragraph(str(h), header_cell_style) for h in headers])

        # Limit PDF rows if too massive (PDFs are summaries)
        pdf_rows = rows[:200]
        for row in pdf_rows:
            table_data.append([Paragraph(str(c if c is not None else ""), cell_style) for c in row])

        if table_data:
            t = Table(table_data, repeatRows=1)
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D3D3D3")),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F9")]),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            elements.append(t)

        if len(rows) > 200:
            overflow_style = ParagraphStyle("Overflow", parent=styles["Normal"], fontSize=8, textColor=colors.gray)
            elements.append(Spacer(1, 0.1 * inch))
            elements.append(
                Paragraph(f"* Note: Showing first 200 of {len(rows)} records. Use XLSX or CSV for full data.", overflow_style)
            )

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    # ---------------------------------------------------------
    # Report Data Retrieval & Generation
    # ---------------------------------------------------------

    def generate_report(
        self,
        db: Session,
        business_id: int,
        request: GenerateReportRequest,
    ) -> Tuple[bytes, str, str, Dict[str, Any]]:
        """
        Fetch data for report_type and render into the requested file_format.
        Returns (file_bytes, file_name, mime_type, metadata_json).
        """
        biz = db.execute(select(Business).where(Business.id == business_id)).scalar_one_or_none()
        biz_name = biz.name if biz else f"Business #{business_id}"

        # Fetch headers and rows depending on report_type
        if request.report_type == "sales_summary":
            title, headers, rows, meta = self._fetch_sales_summary(db, business_id, request)
        elif request.report_type == "financial_summary":
            title, headers, rows, meta = self._fetch_financial_summary(db, business_id, request)
        elif request.report_type == "inventory_status":
            title, headers, rows, meta = self._fetch_inventory_status(db, business_id, request)
        elif request.report_type == "forecast_export":
            title, headers, rows, meta = self._fetch_forecast_export(db, business_id, request)
        elif request.report_type == "promotion_summary":
            title, headers, rows, meta = self._fetch_promotion_summary(db, business_id, request)
        elif request.report_type == "action_center_summary":
            title, headers, rows, meta = self._fetch_action_center_summary(db, business_id, request)
        else:
            raise ValidationException(
                errors={"report_type": f"Unsupported report type: {request.report_type}"},
                error_code=ErrorCodeRegistry.ART_UNSUPPORTED_TYPE,
            )

        report_title = request.title or title
        timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_name = f"{request.report_type}_{timestamp_str}.{request.file_format}"
        mime_type = MIME_TYPES[request.file_format]

        subtitle = f"Business: {biz_name} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        if request.start_date or request.end_date:
            subtitle += f" | Period: {request.start_date or 'Start'} to {request.end_date or 'End'}"

        if request.file_format == "csv":
            file_bytes = self._export_csv(headers, rows)
        elif request.file_format == "xlsx":
            file_bytes = self._export_xlsx(request.report_type, headers, rows)
        elif request.file_format == "pdf":
            file_bytes = self._export_pdf(
                title=report_title,
                subtitle=subtitle,
                headers=headers,
                rows=rows,
                landscape_orientation=len(headers) > 5,
            )
        else:
            raise ValidationException(
                errors={"file_format": f"Unsupported file format: {request.file_format}"},
                error_code=ErrorCodeRegistry.ART_UNSUPPORTED_FORMAT,
            )

        metadata = {
            "business_id": business_id,
            "report_type": request.report_type,
            "row_count": len(rows),
            "generated_at": datetime.now().isoformat(),
            **meta,
        }
        return file_bytes, file_name, mime_type, metadata

    # ---------------------------------------------------------
    # Plan Artifact Generation
    # ---------------------------------------------------------

    def generate_plan_artifact(
        self,
        db: Session,
        business_id: int,
        request: GeneratePlanArtifactRequest,
    ) -> Tuple[bytes, str, str, Dict[str, Any]]:
        """
        Generate downloadable artifact for a PromotionPlan or structured business plan.
        """
        plan = db.execute(
            select(PromotionPlan).where(
                PromotionPlan.id == request.source_plan_id,
                PromotionPlan.business_id == business_id,
            )
        ).scalar_one_or_none()

        if not plan:
            raise NotFoundException(
                message=f"Promotion plan #{request.source_plan_id} not found",
                error_code=ErrorCodeRegistry.PROMO_PLAN_NOT_FOUND,
            )

        plan_title = request.title or f"Promotion Plan: {plan.title}"
        timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_name = f"promotion_plan_{plan.id}_{timestamp_str}.{request.file_format}"
        mime_type = MIME_TYPES[request.file_format]

        # Extract items table
        headers = ["Product / Item", "Discount Type", "Discount Value", "Final Price", "Est. Cost", "Est. Margin", "Margin %"]
        rows = []
        for item in (plan.items or []):
            p_name = item.get("product_name") or f"Product #{item.get('product_id', '')}"
            d_type = item.get("discount_type", "")
            d_val = f"{item.get('discount_value', 0)}%" if d_type == "percentage" else str(item.get("discount_value", 0))
            f_price = str(item.get("discounted_price") or item.get("bundle_price") or item.get("original_price") or 0)
            e_cost = str(item.get("estimated_cost") or "N/A")
            e_margin = str(item.get("estimated_margin") or "N/A")
            m_pct = f"{item.get('margin_pct', 'N/A')}%" if item.get("margin_pct") is not None else "N/A"
            rows.append([p_name, d_type, d_val, f_price, e_cost, e_margin, m_pct])

        summary_text = (
            f"Status: {plan.status.upper()} | Type: {plan.promotion_type} | Dates: {plan.start_date} to {plan.end_date}\n"
            f"Original Revenue: {plan.total_original_revenue or 0:,.2f} | Discounted Revenue: {plan.total_discounted_revenue or 0:,.2f} | "
            f"Est. Margin: {plan.total_estimated_margin or 0:,.2f} ({plan.avg_margin_pct or 0:.2f}%)\n"
            f"Financially Validated: {'Yes' if plan.is_financially_validated else 'No'}"
        )
        if plan.rationale:
            summary_text += f"\nRationale: {plan.rationale}"

        if request.file_format == "csv":
            file_bytes = self._export_csv(headers, rows)
        elif request.file_format == "xlsx":
            summary_rows = [
                ["SUMMARY", "", "", "", "", "", ""],
                ["Original Revenue", str(plan.total_original_revenue or 0), "", "", "", "", ""],
                ["Discounted Revenue", str(plan.total_discounted_revenue or 0), "", "", "", "", ""],
                ["Total Estimated Cost", str(plan.total_estimated_cost or 0), "", "", "", "", ""],
                ["Total Estimated Margin", str(plan.total_estimated_margin or 0), "", "", "", "", ""],
                ["Average Margin %", f"{plan.avg_margin_pct or 0}%", "", "", "", "", ""],
            ]
            file_bytes = self._export_xlsx("Plan Items", headers, rows, summary_rows)
        elif request.file_format == "pdf":
            file_bytes = self._export_pdf(
                title=plan_title,
                subtitle=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
                headers=headers,
                rows=rows,
                summary_text=summary_text,
            )
        else:
            raise ValidationException(
                errors={"file_format": f"Unsupported format: {request.file_format}"},
                error_code=ErrorCodeRegistry.ART_UNSUPPORTED_FORMAT,
            )

        metadata = {
            "source_plan_id": plan.id,
            "plan_title": plan.title,
            "item_count": len(rows),
            "status": plan.status,
            "avg_margin_pct": float(plan.avg_margin_pct or 0),
        }
        return file_bytes, file_name, mime_type, metadata

    # ---------------------------------------------------------
    # Report Data Retrieval Implementations
    # ---------------------------------------------------------

    def _fetch_sales_summary(
        self,
        db: Session,
        business_id: int,
        request: GenerateReportRequest,
    ) -> Tuple[str, List[str], List[List[Any]], Dict[str, Any]]:
        query = (
            select(
                SalesReportObservation.period_start,
                SalesReportObservation.period_end,
                Branch.name.label("branch_name"),
                Product.name.label("product_name"),
                func.sum(SalesReportObservation.sold_qty).label("total_qty"),
                func.sum(SalesReportObservation.recorded_revenue).label("total_sales"),
            )
            .join(Product, SalesReportObservation.product_id == Product.id)
            .join(Branch, SalesReportObservation.branch_id == Branch.id)
            .where(
                SalesReportObservation.business_id == business_id,
                SalesReportObservation.record_status == "active",
            )
        )
        if request.branch_id:
            query = query.where(SalesReportObservation.branch_id == request.branch_id)
        if request.start_date:
            query = query.where(SalesReportObservation.period_start >= request.start_date)
        if request.end_date:
            query = query.where(SalesReportObservation.period_end <= request.end_date)

        query = query.group_by(
            SalesReportObservation.period_start,
            SalesReportObservation.period_end,
            Branch.name,
            Product.name,
        ).order_by(SalesReportObservation.period_start.desc())

        results = db.execute(query).fetchall()
        headers = ["Period Start", "Period End", "Branch", "Product", "Quantity Sold", "Recorded Revenue"]
        rows = [
            [
                r.period_start.strftime("%Y-%m-%d") if r.period_start else "",
                r.period_end.strftime("%Y-%m-%d") if r.period_end else "",
                r.branch_name,
                r.product_name,
                float(r.total_qty or 0),
                float(r.total_sales or 0),
            ]
            for r in results
        ]
        return "Sales Summary Report", headers, rows, {"total_records": len(rows)}

    def _fetch_financial_summary(
        self,
        db: Session,
        business_id: int,
        request: GenerateReportRequest,
    ) -> Tuple[str, List[str], List[List[Any]], Dict[str, Any]]:
        # Fetch revenues from observations
        obs_query = (
            select(
                func.coalesce(func.sum(SalesReportObservation.recorded_revenue), 0).label("revenue"),
            )
            .where(
                SalesReportObservation.business_id == business_id,
                SalesReportObservation.record_status == "active",
            )
        )
        if request.start_date:
            obs_query = obs_query.where(SalesReportObservation.period_start >= request.start_date)
        if request.end_date:
            obs_query = obs_query.where(SalesReportObservation.period_end <= request.end_date)

        obs_res = db.execute(obs_query).fetchone()
        revenue = float(obs_res.revenue if obs_res else 0)

        # Fetch expenses
        exp_query = (
            select(
                Expense.category,
                func.sum(Expense.amount).label("total_amount"),
            )
            .where(Expense.business_id == business_id)
            .group_by(Expense.category)
        )
        if request.start_date:
            exp_query = exp_query.where(Expense.expense_date >= request.start_date)
        if request.end_date:
            exp_query = exp_query.where(Expense.expense_date <= request.end_date)

        exp_results = db.execute(exp_query).fetchall()
        total_expenses = sum(float(e.total_amount or 0) for e in exp_results)
        net_profit = revenue - total_expenses

        headers = ["Financial Metric / Category", "Amount"]
        rows = [
            ["Total Recorded Revenue", revenue],
        ]
        for e in exp_results:
            rows.append([f"Expense: {e.category}", float(e.total_amount or 0)])

        rows.append(["Total Operating Expenses", total_expenses])
        rows.append(["Estimated Operating Profit", net_profit])

        return "Financial Summary Report", headers, rows, {
            "revenue": revenue,
            "total_expenses": total_expenses,
            "net_profit": net_profit,
        }

    def _fetch_inventory_status(
        self,
        db: Session,
        business_id: int,
        request: GenerateReportRequest,
    ) -> Tuple[str, List[str], List[List[Any]], Dict[str, Any]]:
        query = (
            select(
                Product.sku,
                Product.name,
                Product.category,
                Product.base_unit,
                Product.status,
            )
            .where(Product.business_id == business_id)
            .order_by(Product.name.asc())
        )
        results = db.execute(query).fetchall()
        headers = ["SKU", "Product Name", "Category", "Unit", "Status"]
        rows = [
            [
                r.sku or "",
                r.name,
                r.category or "",
                r.base_unit or "",
                r.status,
            ]
            for r in results
        ]
        return "Product & Inventory Status Report", headers, rows, {"product_count": len(rows)}

    def _fetch_forecast_export(
        self,
        db: Session,
        business_id: int,
        request: GenerateReportRequest,
    ) -> Tuple[str, List[str], List[List[Any]], Dict[str, Any]]:
        # Fetch latest forecast run for business
        run_query = (
            select(ForecastRun)
            .where(ForecastRun.business_id == business_id)
            .order_by(ForecastRun.created_at.desc())
        )
        if request.branch_id:
            run_query = run_query.where(ForecastRun.branch_id == request.branch_id)
        latest_run = db.execute(run_query).scalars().first()

        headers = ["Forecast Date", "Predicted Demand", "Weekday", "Method"]
        rows = []
        meta = {}
        if latest_run:
            meta["run_id"] = latest_run.id
            meta["method"] = latest_run.method
            meta["mae"] = float(latest_run.mae or 0)
            res_query = (
                select(ForecastResult)
                .where(ForecastResult.forecast_run_id == latest_run.id)
                .order_by(ForecastResult.forecast_date.asc())
            )
            for res in db.execute(res_query).scalars().all():
                rows.append([
                    res.forecast_date.strftime("%Y-%m-%d"),
                    float(res.predicted_qty),
                    res.weekday if res.weekday is not None else "",
                    latest_run.method or "SMA",
                ])

        return "Demand Forecast Export", headers, rows, meta

    def _fetch_promotion_summary(
        self,
        db: Session,
        business_id: int,
        request: GenerateReportRequest,
    ) -> Tuple[str, List[str], List[List[Any]], Dict[str, Any]]:
        query = (
            select(PromotionPlan)
            .where(PromotionPlan.business_id == business_id)
            .order_by(PromotionPlan.created_at.desc())
        )
        plans = db.execute(query).scalars().all()
        headers = [
            "Plan ID", "Title", "Type", "Status", "Start Date", "End Date",
            "Original Rev", "Discounted Rev", "Est Margin", "Avg Margin %", "Validated"
        ]
        rows = [
            [
                p.id,
                p.title,
                p.promotion_type,
                p.status,
                p.start_date.strftime("%Y-%m-%d") if p.start_date else "",
                p.end_date.strftime("%Y-%m-%d") if p.end_date else "",
                float(p.total_original_revenue or 0),
                float(p.total_discounted_revenue or 0),
                float(p.total_estimated_margin or 0),
                f"{p.avg_margin_pct or 0}%",
                "Yes" if p.is_financially_validated else "No",
            ]
            for p in plans
        ]
        return "Promotion Plans Summary Report", headers, rows, {"total_plans": len(plans)}

    def _fetch_action_center_summary(
        self,
        db: Session,
        business_id: int,
        request: GenerateReportRequest,
    ) -> Tuple[str, List[str], List[List[Any]], Dict[str, Any]]:
        query = (
            select(BusinessAction)
            .where(BusinessAction.business_id == business_id)
            .order_by(BusinessAction.priority_score.desc())
        )
        actions = db.execute(query).scalars().all()
        headers = [
            "Action ID", "Title", "Category", "Action Type", "Urgency", "Impact",
            "Priority Score", "Status", "Rule"
        ]
        rows = [
            [
                a.id,
                a.title,
                a.category,
                a.action_type,
                a.urgency,
                a.expected_impact,
                float(a.priority_score or 0),
                a.status,
                a.rule_label,
            ]
            for a in actions
        ]
        return "Business Action Center Summary Report", headers, rows, {"total_actions": len(actions)}
